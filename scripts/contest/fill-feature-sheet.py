# 공모전 「기능 설명서」 공식 양식(pptx)을 채운다. 2026-09-20 작성 · 2026-09-21 내용 갱신(순례 기록 축·6개 국어 완비·분산 세 축).
# 사용: PYTHONIOENCODING=utf-8 python scripts/contest/fill-feature-sheet.py <양식.pptx> <스크린샷 폴더> <출력.pptx>
# 양식은 9장 고정. 5장(기능 흐름도)은 해시태그 수만큼 복제한다. 안내문 상자(빨간 박스)는 지운다.
import copy, sys, io
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

SRC, SHOTS, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
BLACK = RGBColor(0, 0, 0)
SITES = '206'  # DB 실측 2026-09-20 (소개글 206×6개 국어 · 오디오 가이드 108×6)


def set_text(cell_or_shape, lines, size=14, bold_first=False, align=None):
    """셀(또는 도형)의 안내문을 지우고 줄 단위로 채운다. 글꼴은 테마(+mn-lt)를 그대로 물려받는다."""
    tf = cell_or_shape.text_frame
    tf.word_wrap = True
    first = tf.paragraphs[0]
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, line in enumerate(lines):
        p = first if i == 0 else tf.add_paragraph()
        if align is not None:
            p.alignment = align
        # 「**굵게**」 표시를 굵은 run 으로 — 읽는 사람이 한눈에 요점을 잡게(사장님 지시 2026-09-21)
        for j, piece in enumerate(line.split('**')):
            if piece == '':
                continue
            run = p.add_run()
            run.text = piece
            run.font.size = Pt(size)
            run.font.color.rgb = BLACK
            run.font.bold = (bold_first and i == 0) or (j % 2 == 1)


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def table_of(slide, name=None):
    for sh in slide.shapes:
        if sh.has_table and (name is None or sh.name == name):
            return sh
    raise KeyError(name)


def cell_box(shape, r, c):
    """표 안 (r,c) 셀의 절대 위치·크기(EMU). 사진은 표 셀에 못 넣으므로 그 위에 겹쳐 놓는다."""
    t = shape.table
    left = shape.left + sum(col.width for col in list(t.columns)[:c])
    top = shape.top + sum(row.height for row in list(t.rows)[:r])
    return left, top, t.columns[c].width, t.rows[r].height


def put_picture(slide, path, left, top, w, h, pad=Emu(60000)):
    """상자 안에 비율 유지로 맞춰 가운데 놓는다."""
    iw, ih = Image.open(path).size
    bw, bh = w - 2 * pad, h - 2 * pad
    scale = min(bw / iw, bh / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(path, left + (w - pw) // 2, top + (h - ph) // 2, pw, ph)


def duplicate_slide(prs, src):
    """python-pptx 에 복제가 없어 도형 XML 을 통째로 복사한다. 그림은 뒤에 새로 넣으므로 표·텍스트만 복사한다."""
    dst = prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes):
        remove_shape(sh)
    for sh in src.shapes:
        dst.shapes._spTree.append(copy.deepcopy(sh._element))
    return dst


def move_slide(prs, slide, new_index):
    lst = prs.slides._sldIdLst
    el = [e for e in lst if e.rId == prs.part.relate_to(slide.part, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide')][0]
    lst.remove(el)
    lst.insert(new_index, el)


prs = Presentation(SRC)
S = list(prs.slides)

# ── 1. 표지 ─────────────────────────────────────────────
t = S[0].shapes[0].table
set_text(t.cell(0, 1), ['visitholykorea'], 18)
set_text(t.cell(1, 1), ['VisitHolyKorea — 길 위에서 나를 만나는 성지순례'], 18)

# ── 2. 서비스 소개 ───────────────────────────────────────
t = table_of(S[1]).table
rows = [
    (['**VisitHolyKorea** — 길 위에서 나를 만나는 성지순례  (visitholykorea.com)'], 18),
    (['**반응형 웹** — 휴대폰(하단 탭)과 PC(상단 메뉴)를 **한 코드**로, 1024px 기준 자동 전환',
      '홈 화면 설치(PWA)로 앱처럼 실행 · 같은 코드로 Android/iOS 앱(Capacitor) 빌드'], 15),
    (['① 붐비는 명소 대신 **조용한 하루**를 원하는 자유여행객·시니어 (이용자 조사 응답자 **85%가 50대 이상**)',
      '② 한국 천주교 **성지 순례객**과 **2027 서울 세계청년대회(WYD)** 앞뒤로 방한하는 외국인 순례객 (**6개 국어**)',
      '③ 종교와 무관하게 걷기·사색·역사 여행에 관심 있는 사람'], 15),
    ([f'전국 천주교 성지 **{SITES}곳**을 "오늘 마음·출발지·시간"에 맞춰 골라 주고,',
      '한국관광공사 TourAPI **실시간 데이터**(관광지 집중률 예측·주변 관광지·걷기길)로 **하루 일정**을 짜 주며,',
      '현장에서는 **6개 국어 성지 이야기와 오디오 가이드**로 안내하고, 다녀온 **순례 기록**을 함께 남기는 순례 여행 서비스'], 15),
    (['**지정과제 2번 — 유명 관광지 쏠림(오버투어리즘) 문제 해결**'], 17),
    ([f'• 성지 {SITES}곳 중 **140곳(68%)이 비수도권**(2026-09-21 DB 실측) — 목적지 목록 자체가 분산을 돕는다',
      '• 순례자는 **목적지가 정해진 여행자** — "명동 대신 어디"가 아니라 "명동 성당 다음엔 어디"로 물으면 동선이 바뀐다',
      '• 관광공사 **집중률 예측을 실시간**으로 엮어 붐비는 날·곳을 피해 **같은 지역의 조용한 성지**와 볼거리를 제안',
      f'• 유명 성지 열 곳에 몰리던 순례를 **공간({SITES}곳)·시간(30일 예측)·체류(주변 식당·걷기길)**로 펼친다'], 14),
]
for i, (lines, size) in enumerate(rows):
    set_text(t.cell(i, 1), lines, size)

# ── 3. 기획 방향 · 해시태그 목록 ────────────────────────
t = table_of(S[2]).table
set_text(t.cell(0, 1), [
    f'**1) 문제와 범위** — 성지 순례는 유명 성지 열 곳·주말·수도권에 몰린다. 일반 관광객이 아니라 **순례자의 동선을 수도권 밖·붐비는 날 밖으로** 넓히는 것이 목표. 성지 {SITES}곳 중 **140곳(68%)이 비수도권**.',
    '**2) 공간 분산(어디로)** — 관광공사 「관광지 집중률 예측」을 실시간으로 받아(저장 안 함) 성지마다 **「주변 밀집도 상·중·하」**(숫자 대신 세 단계, 50대 이상도 한눈에). 「오늘의 성지 일정」은 마음·출발지·시간(질문 3개)에 맞는 가까운 후보 3곳에 **「인근이 조용해요」 표시로 밀집도 낮은 성지를 고르게** 하고, 오후가 붐빌 예정이면 같은 지역의 다른 관광지를 제안.',
    '**3) 체류(얼마나 오래)** — 고른 성지를 축으로 **오전 성지 → 점심(주변 식당) → 오후(주변 관광지·걷기길)** 하루 일정을 TourAPI 실시간 조회로(응답 저장 안 함). 순례 코스 11개(6개 국어)로 1곳 방문을 여러 곳·여러 날로.',
    f'**4) 현장 콘텐츠** — 성지 {SITES}곳 소개글과 108곳 지점별 오디오 가이드를 **6개 국어로 직접 집필**, 성지 DB 안에서만 답하는 **AI 가이드** — 지방의 덜 알려진 성지도 "갈 만한 곳"으로.',
    '**5) 시간 분산(언제)과 측정** — 집중률 예측 30일치를 **「한적한 날」** 카드로 펼쳐 "주말 말고 이 날에". **「순례 기록」**을 서비스 중심에 두고 저장 때 **「그날 붐볐나요?」**를 받아 예측을 실측으로 보정할 재료를 쌓고(보정 산식은 발전계획), 기록의 성지 → 지역으로 **비수도권 방문 비율을 측정**.',
], 13)
set_text(t.cell(1, 1), [
    '#오늘의성지일정  #오디오가이드  #순례가이드미카엘  #순례기록',
], 18, bold_first=True)

# ── 4. 해시태그 연계 표 ─────────────────────────────────
HASHTAGS = [
    ('#오늘의성지일정', '마음 상태로 고르는 하루 일정 — 오버투어리즘 대응의 중심',
     ['• **질문 3개**(마음 · 출발지 · 시간) → 반경 안 **후보 성지 3곳** → 하나를 고르면 **오전 성지 → 점심 → 오후 관광지** 하루 일정',
      '• TourAPI **관광지 집중률 예측**으로 「인근이 조용해요」 후보 표시와 「인근 지역이 조용해요/보통이에요/붐벼요」, locationBasedList2 로 주변 식당·관광지 **실시간** 조회. 오후가 붐빌 예정이면 성지에 더 머무르기 또는 **같은 반경의 다른 관광지로 바꾸기** 제안',
      '• 시간 답(반나절·하루·1박2일) = 출발지 반경 **20·60·180km** — 이동이 체류보다 길어지지 않게. 성지 상세의 **「주변 밀집도 상·중·하」**와 30일 예측 중 조용한 날짜를 고른 **「한적한 날」** 카드 — 붐비는 날을 피하는 시간 분산']),
    ('#오디오가이드', '성지 이야기 · 지점별 오디오 가이드 — 6개 국어',
     [f'• 성지 **{SITES}곳** 소개글(위치·역사·건축·주변 이야기 3문단)을 **6개 국어**(한·영·서·프·포·이)로 직접 집필 — 2026-09-21 실측 {SITES} × 6 완비',
      '• 현장 지점 원고(여는 말 → 지점 3~6곳 → 맺음말) **108곳 × 6개 국어**. 기기 TTS 낭독·속도 조절·챕터 이동. 원고 없는 성지는 소개·역사 낭독으로 대체',
      '• 성지 이야기에 **참고한 출처 링크** 표시(교구 공식 자료·굿뉴스 성지 안내). 두루누비 걷기길·둘러볼 곳(관광공사 실시간)과 나란히. 지구본 버튼 하나로 **화면·이야기·가이드가 함께 언어 전환**']),
    ('#순례가이드미카엘', '성지 DB 범위 안에서만 답하는 AI 가이드',
     ['• 어느 화면에서나 상단 「미카엘」. 질문과 관련된 **성지 레코드(주소·전화·미사·설명)를 컨텍스트**로 넣고 "자료에 없으면 모른다고 답하라"를 시스템 프롬프트에 고정 — **지어내지 않는다**',
      '• **Claude 우선, 실패 시 Gemini** 이중 구조. Gemini 모델 목록은 GitHub Actions 가 매일 갱신. 키는 Supabase Edge Function 에만(브라우저 비노출)',
      '• 답변에 **「참고한 성지」** 표시. 로그인하면 대화가 계정에 저장·이어짐(「대화 지우기」로 삭제). **6개 국어 질문**에 같은 DB 근거로 답한다']),
    ('#순례기록', '함께 만드는 순례 기록 — 내 한 줄이 다음 순례자의 안내가 된다',
     ['• 성지 상세 「순례 기록 남기기」: 방문일 · 한 줄 메모 · **사진 최대 3장**. 기록은 그 성지 페이지 아래에 **다른 순례자의 후기로 쌓여** 「순례 기록 (N)」 — 운영자가 아니라 **모두가 함께 만드는 순례 안내**. 익명 공개·신고',
      '• 하단 탭 가운데 「기록」에서 내 기록·즐겨찾기 모아 보기·수정·삭제. 순례 코스는 기록 기준 「n화까지 기록」 진행. 지금은 성지당 1개 → **재방문도 기록**하도록 확장',
      '• 저장 때 **「그날 붐볐나요?」**(한적·보통·붐빔) 한 칸 — 관광공사 **예측을 실측으로 보정**하는 재료. 기록은 우리 데이터라 **비수도권 방문 비율을 측정**할 수 있다']),
]
sh4 = table_of(S[3])
t = sh4.table
while len(t.rows) < 1 + len(HASHTAGS):
    t._tbl.append(copy.deepcopy(t._tbl.tr_lst[-1]))
for i, (tag, feat, detail) in enumerate(HASHTAGS, 1):
    t.rows[i].height = Emu(int(4543492 / len(HASHTAGS)))
    set_text(t.cell(i, 0), [tag], 14, bold_first=True)
    set_text(t.cell(i, 1), [feat], 12)
    set_text(t.cell(i, 2), detail, 10.5)
for sh in list(S[3].shapes):
    if sh.shape_type == 1:  # 빨간 안내 상자
        remove_shape(sh)

# ── 5. 기능 흐름도 — 해시태그마다 한 장 ──────────────────
FLOWS = {
    '#오늘의성지일정': [
        ('home-390.png', '홈 입구 카드 「오늘의 성지 일정」'),
        ('compass-q1-390.png', '지금 마음 → 출발지(현재 위치·시도) → 낼 수 있는 시간'),
        ('compass-result-390.png', '반경 안 후보 3곳 — 가장 가까움·인근이 조용함·소개가 자세함 표시'),
        ('detail-quietdays-390.png', '성지 상세 「한적한 날」 — 30일 예측 중 조용한 날짜. 일정 결과는 오전 성지 → 점심 → 오후'),
    ],
    '#오디오가이드': [
        ('detail-top-390.png', '성지 상세 — 사진·분류·주변 밀집도 상/중/하·태그'),
        ('detail-docent-390.png', '「오디오 가이드」 여는 말 → 지점 → 맺음말, 속도 조절, 눈여겨보기'),
        ('detail-docent-en-390.png', '지구본 버튼으로 언어를 바꾸면 가이드도 그 언어로 (영어 예)'),
        ('detail-nearby-390.png', '「둘러볼 곳」 — 관광공사 실시간 식당·볼거리로 체류 연장'),
    ],
    '#순례가이드미카엘': [
        ('michael-390.png', '질문 — 공세리 성지 위치와 전화번호. 「참고한 성지」 표시'),
        ('michael-en-390.png', '외국어 질문에도 같은 DB 근거로 답한다 (영어 예)'),
        ('michael-mass-390.png', 'DB 에 없는 정보는 지어내지 않고 사무실 확인을 안내'),
        ('michael-empty-390.png', '어느 화면에서나 상단 「미카엘」 버튼 — 로그인하면 대화가 이어진다'),
    ],
    '#순례기록': [
        ('detail-record-390.png', '성지 상세 「순례 기록」 — 다녀온 사람들의 후기가 쌓이는 자리, 사진 최대 3장'),
        ('crowded-q-390.png', '기록 남기기의 「그날 붐볐나요?」 — 예측을 검증할 실측 한 칸'),
        ('login-390.png', '카카오·네이버·구글 간편 로그인, 이메일 가입(확인 메일은 화면 언어로)'),
        ('route-detail-390.png', '순례 코스 상세 — 기록 기준 「n화까지 기록」 진행'),
    ],
}
flow_src = S[4]
flow_slides = [flow_src]
for _ in range(len(HASHTAGS) - 1):
    flow_slides.append(duplicate_slide(prs, flow_src))
for k, sl in enumerate(flow_slides[1:], 1):
    move_slide(prs, sl, 4 + k)

for sl, (tag, feat, detail) in zip(flow_slides, HASHTAGS):
    for sh in list(sl.shapes):
        if sh.shape_type == 1:
            remove_shape(sh)
    t5 = table_of(sl, '표 5').table
    n = HASHTAGS.index((tag, feat, detail)) + 1
    for r, head in enumerate(['해시태그', '연계기능', '기능설명']):
        set_text(t5.cell(r, 0), [f'{head}{n}'], 14, bold_first=True, align=PP_ALIGN.CENTER)
        for para in t5.cell(r, 0).text_frame.paragraphs:
            for run in para.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    set_text(t5.cell(0, 1), [tag], 14, bold_first=True)
    set_text(t5.cell(1, 1), [feat], 13)
    set_text(t5.cell(2, 1), [' '.join(d.lstrip('• ') for d in detail)], 9.5)
    sh4f = table_of(sl, '표 4')
    t4 = sh4f.table
    t4.rows[1].height = Emu(2900000)
    t4.rows[2].height = Emu(560000)
    for c, (img, step) in enumerate(FLOWS[tag]):
        set_text(t4.cell(1, c), [''], 8)
        left, top, w, h = cell_box(sh4f, 1, c)
        put_picture(sl, f'{SHOTS}/{img}', left, top, w, h)
        set_text(t4.cell(2, c), [f'{c + 1}. {step}'], 10.5)

# ── 6. 서비스 관련 이미지 ───────────────────────────────
sl = S[5]
sh6 = table_of(sl)
t = sh6.table
set_text(t.cell(0, 1), [''], 8)
set_text(t.cell(1, 1), [''], 8)
left, top, w, h = cell_box(sh6, 0, 1)
put_picture(sl, 'public/logo.png', left, top, w, h, pad=Emu(300000))
left, top, w, h = cell_box(sh6, 1, 1)
# 왼쪽 55% 는 PC 화면(반응형 강조), 오른쪽은 휴대폰 3장 — 같은 코드가 두 폭에서 어떻게 보이는지
pc_w = int(w * 0.55)
put_picture(sl, f'{SHOTS}/detail-desktop-1440.png', left, top, pc_w, h, pad=Emu(40000))
imgs = ['detail-top-390.png', 'compass-plan-390.png', 'michael-390.png']
cw = (w - pc_w) // len(imgs)
for i, img in enumerate(imgs):
    put_picture(sl, f'{SHOTS}/{img}', left + pc_w + cw * i, top, cw, h, pad=Emu(40000))

# ── 7. 한국관광공사 OpenAPI ─────────────────────────────
APIS = [
    ('관광지 집중률 예측 서비스 (TatsCnctrRateService) — 핵심 지표',
     '성지 소재 시·군·구의 집중률 예측(오늘부터 30일)을 받아 **「주변 밀집도 상·중·하」**(성지 상세), **「인근이 조용해요」**(일정 후보), **「한적한 날」**(30일 중 조용한 날짜)에 쓴다. 메모리에만 잠깐 두고 **저장하지 않는다**.'),
    ('국문 관광정보 서비스 (KorService2)',
     'locationBasedList2 로 성지 반경 5km 의 **관광지·식당**을 받아 성지 상세 「둘러볼 곳」과 하루 일정의 점심·오후(붐빌 예정이면 같은 풀의 다음 후보로 대안)를 채우고, ldongCode2 로 집중률 조회에 필요한 시·군·구 코드를 받는다. **실시간 호출**(응답은 화면 메모리에서만 잠시 유지), DB·서비스워커에 저장하지 않음.'),
    ('영문 관광정보 서비스 (EngService2)',
     '앱 언어가 영어·포르투갈어·이탈리아어일 때 주변 관광지를 **영문 서비스**의 locationBasedList2 로 조회 — 2027 WYD 외국인 순례객이 **읽을 수 있는 주변 정보**. 결과가 비면 국문으로 폴백.'),
    ('서어·불어 관광정보 서비스 (SpnService2 · FreService2)',
     '앱 언어가 스페인어·프랑스어일 때 같은 조회를 **그 언어 서비스**로 보낸다(WYD 참가국이 많은 중남미·프랑스어권). 결과가 비면 국문으로 폴백.'),
    ('두루누비 걷기길 서비스 (Durunubi)',
     'courseList(DNWW)로 성지와 같은 시·군·구의 걷기길을 받아 성지 상세 **「이 근처 걷기길」**로 표시 — 순례 후 도보 체류를 늘린다.'),
]
t = table_of(S[6]).table
for i, (name, desc) in enumerate(APIS):
    set_text(t.cell(2 * i, 2), [name], 12, bold_first=True)
    set_text(t.cell(2 * i + 1, 2), [desc], 10.5)
for i in range(len(APIS), len(t.rows) // 2):  # 안 쓰는 행은 비운다 — 없는 API 를 적지 않는다
    set_text(t.cell(2 * i, 2), [''], 10)
    set_text(t.cell(2 * i + 1, 2), [''], 10)

# ── 8. 기타 데이터 ─────────────────────────────────────
OTHERS = [
    (f'자체 구축 성지 DB (holy_sites · holy_site_translations · docent_scripts, {SITES}곳)',
     f'각 교구 홈페이지·굿뉴스 성지 안내·가톨릭 언론의 공개 자료와 현장에서 수집한 주보·안내문을 대조해 **직접 쓴** 성지 {SITES}곳의 주소·연락처·미사 시간·역사. 성지마다 **6개 국어 소개글({SITES}×6)**과 **지점별 오디오 가이드(108×6)**. 대전교구 성지는 교구 홍보국 제공 자료 참고. 성지 이야기에 **출처 링크** 표시. 대표 사진은 직접 촬영·대전교구 제공·위키미디어 커먼즈(CC)·국가유산청(공공누리) — 라이선스가 요구하는 것은 화면에 출처 표기. Supabase(PostgreSQL).'),
    ('한국천주교주소록 (CBCK, catholic_directory 5,918건)',
     '한국천주교중앙협의회 공개 주소록 → 성지 상세·시도 화면의 **「주변 성당·공소」**(가까운 미사). 외국어 화면은 자동 로마자.'),
    ('순례 기록 (pilgrimage_stamps · 이용자 생성)',
     '방문일 · 한 줄 · 사진 3장 · **그날 붐볐나요**. 우리 데이터라 분석 가능 — **비수도권 방문 비율 측정, 집중률 예측 보정**의 재료.'),
    ('AI API — Anthropic Claude, Google Gemini',
     '「순례 가이드 미카엘」 답변 생성. 성지 DB 레코드를 컨텍스트로 넣고 **범위 밖은 모른다고** 답하도록 제한. Supabase Edge Function 에서만 호출(키 비노출).'),
]
sl = S[7]
t = table_of(sl).table
for i, (name, desc) in enumerate(OTHERS):
    if 2 * i + 1 >= len(t.rows):
        break
    set_text(t.cell(2 * i, 2), [name], 12, bold_first=True)
    set_text(t.cell(2 * i + 1, 2), [desc], 10.5)
for sh in list(sl.shapes):
    if sh.shape_type == 1:
        remove_shape(sh)

# ── 9. 차별성 · 발전계획 ───────────────────────────────
t = table_of(S[8]).table
set_text(t.cell(0, 1), [
    '• **웹 구현** — 반응형(휴대폰 하단 탭·PC 상단 메뉴를 **한 코드**로, 1024px 자동 전환) · PWA 홈 화면 설치 · **6개 국어** 화면·이야기·가이드·순례 코스 · 큰 글자 모드 · 접근성 정적 검토(WCAG 2.1 기준) 상위 항목 반영 · 카카오·네이버·구글 로그인 · 계정 관리(이름·비밀번호 변경·삭제)',
    f'• **「대체 관광지 추천」이 아니라 「오늘의 하루 일정」 + 「한적한 날」** — 집중률 예측을 실시간으로 조합해 **붐비는 곳과 붐비는 날을 모두 피하면서 그 지역에 머물게** 한다. 성지 **68%(140/{SITES})가 비수도권**이라 분산이 구조적으로 일어난다',
    f'• **콘텐츠를 직접 만들었다** — 성지 {SITES}곳 소개글과 108곳 오디오 가이드를 6개 국어로, 성지 이야기에 출처 링크. 기존 관광 앱이 다루지 않는 **종교 유산 층위**',
    '• **AI 가 지어내지 않는다** — 미카엘은 성지 DB 안에서만 답하고 참고한 성지를 밝힌다',
    '• **함께 만드는 순례 기록** — 다녀온 사람의 한 줄과 사진이 다음 순례자의 안내가 되고, 「그날 붐볐나요?」가 쌓여 **예측을 실측과 대조할 자료**가 된다(보정 산식은 발전계획)',
    '• **관광공사 데이터 원칙 준수** — TourAPI 응답을 DB·서비스워커에 저장하지 않고 실시간 호출(집중률은 화면 메모리에서만 최대 6시간 재사용 — ADR 0002, 코드로 강제)',
], 14)
set_text(t.cell(1, 1), [
    '• **2027 서울 세계청년대회(WYD) — 외국인 순례객 100만 명 시대의 다국어 성지 플랫폼.** 내년 세계청년대회로 외국인 순례객이 대거 방한하지만 **전국 성지를 6개 국어로 안내하는 플랫폼은 없다** — 이 서비스가 그 대안이다. 교구·본당 미사 시간 연동(CBCK 주소록), WYD 공식 일정지(솔뫼·해미) 특별 안내, 단체 순례 일정 공유로 확장',
    '• **분산 측정·개인화** — 기록의 성지 → 지역으로 **비수도권 방문 비율 대시보드**. 기록이 수도권에 몰린 이용자에게 「다음은 갈매못·솔뫼 어떠세요」 제안. 교구별 완주 배지(성지가 적은 교구부터)',
    '• **예측 보정** — 「그날 붐볐나요?」 실측으로 성지별 예측·실제 차이를 계산해 밀집도 표시를 보정(쓸수록 정확). 전례력(순교자 성월·축일 미사)을 「한적한 날」에 함께 표시. 같은 성지 재방문 기록',
    '• **지방 체류** — 지방 성지 **1박 2일** 일정과 **대중교통·순례버스** 접근 정보(차 없는 시니어). 오디오 가이드 나머지 98곳, 대표 사진 전 성지',
    '• **지자체·교구 협업과 스토어 출시** — 시·도 랜딩 페이지로 지역 관광과 협업, Android/iOS 스토어 출시(Capacitor)',
], 14)

prs.save(OUT)
print('saved', OUT, 'slides', len(prs.slides))
