# 사장님 번역 작업지(xlsx, 시트 「지점 원고」)에 data/docent/*.json 의 지점 원고를 채워 넣는다.
# 규칙: 새 파일을 만들지 않고 기존 파일을 고친다. (siteId, 지점) 이 이미 있는 줄은 빈 칸만 채우고
#       사장님이 써 둔 칸은 절대 덮어쓰지 않는다 — 반대로 **사장님 칸이 JSON 과 다르면 JSON 을 사장님 값으로 고친다**
#       (작업지가 번역의 기준, 그 뒤 `npm run docent:load` 로 DB 반영). 번역이 하나도 없는 성지(한국어만)는 넣지 않는다.
# 사용: PYTHONIOENCODING=utf-8 python scripts/docent-xlsx-sync.py [xlsx 경로]
import glob, io, json, os, sys, shutil, datetime
import openpyxl

XLSX = sys.argv[1] if len(sys.argv) > 1 else r'C:\Users\noh hui sun\OneDrive\바탕 화면\번역-작업지-지점원고.xlsx'
COLS = {'번호': 1, '성지': 2, 'siteId': 3, '지점': 4, '제목': 5, 'ko': 6, 'en': 7, 'it': 8, 'fr': 9, 'pt': 10, 'es': 11, '메모': 12}
SFX = {'en': 'En', 'it': 'It', 'fr': 'Fr', 'pt': 'Pt', 'es': 'Es'}

wb = openpyxl.load_workbook(XLSX)
ws = wb['지점 원고']
# 머리글은 2행, 자료는 3행부터
existing = {}
for r in range(3, ws.max_row + 1):
    sid, seq = ws.cell(r, COLS['siteId']).value, ws.cell(r, COLS['지점']).value
    if sid and seq is not None:
        existing[(str(sid).strip(), int(seq))] = r
last = max(existing.values(), default=2)
# 번호는 성지 단위 — 이미 있는 성지는 그 번호를 쓰고, 새 성지는 이어서 매긴다
site_no = {}
for (sid, _), r in existing.items():
    v = ws.cell(r, COLS['번호']).value
    if v is not None and sid not in site_no:
        site_no[sid] = int(v)
next_no = max(site_no.values(), default=0) + 1

added = filled = backfilled = 0
changed_files = set()
for f in sorted(glob.glob('data/docent/*.json')):
    if '_템플릿' in f:
        continue
    d = json.load(io.open(f, encoding='utf-8'))
    sid = d.get('siteId')
    if not sid or not d.get('intro', {}).get('narrationEn'):
        continue  # 번역이 없는 성지는 아직 작업지에 올리지 않는다
    name = d.get('siteName') or d.get('name') or os.path.basename(f)[:-5].replace('-', ' ')
    entries = [(0, d['intro'], '여는 말')] + [(p['seq'], p, p.get('title', '')) for p in d.get('points', [])] + [(99, d['outro'], '맺음말')]
    if sid not in site_no:
        site_no[sid] = next_no
        next_no += 1
    for seq, obj, title in entries:
        key = (sid, seq)
        if key in existing:
            r = existing[key]
        else:
            last += 1
            r = last
            existing[key] = r
            ws.cell(r, COLS['번호'], site_no[sid])
            ws.cell(r, COLS['성지'], name)
            ws.cell(r, COLS['siteId'], sid)
            ws.cell(r, COLS['지점'], seq)
            ws.cell(r, COLS['제목'], title)
            ws.cell(r, COLS['ko'], obj.get('narration', ''))
            ws.cell(r, COLS['메모'], f'Codex 번역 · Claude 검사 {datetime.date.today()}')
            added += 1
        for lang, sfx in SFX.items():
            c = ws.cell(r, COLS[lang])
            v = obj.get('narration' + sfx)
            cell = str(c.value).strip() if c.value and str(c.value).strip() else ''
            if v and not cell:
                c.value = v
                filled += 1
            elif cell and cell != (v or '').strip():
                obj['narration' + sfx] = cell  # 사장님 작업지가 이긴다
                backfilled += 1
                changed_files.add(f)
    if f in changed_files:
        io.open(f, 'w', encoding='utf-8', newline='\n').write(json.dumps(d, ensure_ascii=False, indent=2) + '\n')

shutil.copy(XLSX, XLSX + '.bak')  # 만약을 위해 직전 판 한 벌
wb.save(XLSX)
print(f'줄 추가 {added} · 칸 채움 {filled} · JSON 에 되돌려 쓴 칸 {backfilled}({len(changed_files)}파일) · 총 {last - 2}줄 · 성지 {len(site_no)}곳')
